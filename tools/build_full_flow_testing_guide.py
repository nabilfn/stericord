from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path("docs/SteriCord_Full_Flow_Testing_Guide.docx")
NAVY = "0B2545"; BLUE = "2E74B5"; TEAL = "137C8B"; PALE = "E8EEF5"; LIGHT = "F2F4F7"; RED = "9B1C1C"; GOLD = "7A5A00"; GRAY = "5B6573"; WHITE = "FFFFFF"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc.get_or_add_tcPr(); tcMar = tc.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tc.append(tcMar)
    for side, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = OxmlElement(f"w:{side}"); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); tcMar.append(node)

def fixed_table(table, widths):
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr; tblW = tblPr.first_child_found_in("w:tblW")
    tblW.set(qn("w:w"), str(sum(widths))); tblW.set(qn("w:type"), "dxa")
    ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), "120"); ind.set(qn("w:type"), "dxa"); tblPr.append(ind)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(w)); grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i]/1440); margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcW = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW"); tcW.set(qn("w:w"), str(widths[i])); tcW.set(qn("w:type"), "dxa")

def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"; run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri"); run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None: run.italic = italic

def add_p(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        a,b=text.split(":",1); set_font(p.add_run(a+":"), bold=True); p.add_run(b)
    else: p.add_run(text)
    return p

def bullet(doc, text, level=0):
    return add_p(doc, text, "List Bullet" if level == 0 else "List Bullet 2")

def num(doc, text): return add_p(doc, text, "List Number")

def heading(doc, text, level=1): return doc.add_heading(text, level=level)

def callout(doc, label, text, tone="blue"):
    colors={"blue":(PALE,NAVY),"risk":("FDECEC",RED),"warn":("FFF4D6",GOLD),"ok":("E7F4F1",TEAL)}; fill,ink=colors[tone]
    t=doc.add_table(rows=1, cols=1); fixed_table(t,[9360]); shade(t.cell(0,0),fill)
    p=t.cell(0,0).paragraphs[0]; r=p.add_run(label.upper()+"  "); set_font(r,bold=True,color=ink); p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"; fixed_table(t,widths)
    for i,h in enumerate(headers):
        shade(t.rows[0].cells[i], PALE); p=t.rows[0].cells[i].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(h),bold=True,color=NAVY,size=9)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            p=cells[i].paragraphs[0]; set_font(p.add_run(str(v)),size=9); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
    fixed_table(t,widths); doc.add_paragraph().paragraph_format.space_after=Pt(0); return t

def tc(doc, tid, title, role, pre, steps, expected, evidence, priority="P0"):
    heading(doc, f"{tid} — {title}", 3)
    p=doc.add_paragraph(); set_font(p.add_run(f"{priority}  |  Role: {role}"),bold=True,color=TEAL,size=9)
    add_p(doc, "Preconditions: "+pre)
    for s in steps: num(doc,s)
    add_p(doc, "Expected result: "+expected)
    add_p(doc, "Capture: "+evidence)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1); sec.right_margin=Inches(1); sec.header_distance=Inches(.492); sec.footer_distance=Inches(.492)
styles=doc.styles
normal=styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string("1F2937"); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
for name,size,color,before,after in [("Title",28,NAVY,0,8),("Subtitle",14,GRAY,0,18),("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,14,7),("Heading 3",11,NAVY,10,5)]:
    s=styles[name]; s.font.name="Calibri"; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=name.startswith("Heading"); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
for name in ["List Bullet","List Bullet 2","List Number"]:
    styles[name].font.name="Calibri"; styles[name].font.size=Pt(11); styles[name].paragraph_format.space_after=Pt(4); styles[name].paragraph_format.line_spacing=1.25

# Header/footer
hp=sec.header.paragraphs[0]; hp.text="STERICORD  |  QUALITY ASSURANCE"; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_font(hp.runs[0],size=8,bold=True,color=GRAY)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(fp.add_run("Controlled test guide  •  Version 1.0  •  21 August 2026  •  Page "),size=8,color=GRAY)
fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); fp._p.append(fld)

# Cover
doc.add_paragraph().paragraph_format.space_after=Pt(86)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("SYSTEM TEST MANUAL"),size=10,bold=True,color=TEAL)
p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("SteriCord Full-Flow Testing Guide")
p=doc.add_paragraph(style="Subtitle"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Hospital sterilization traceability, OT safety, evidence, AI assistance, compliance, and resilience")
doc.add_paragraph().paragraph_format.space_after=Pt(28)
table(doc,["Document control","Value"],[
    ("Version","1.0"),("System baseline","Phase 8 complete; repository reviewed 21 Aug 2026"),("Audience","QA, clinical SMEs, CSSD/OT pilot team, security and operations"),("Test level","System, end-to-end, integration, security, recovery and pilot acceptance"),("Primary surfaces","React web, Flutter mobile, Supabase, Edge Functions, reporting service")],[2100,7260])
callout(doc,"Safety principle","AI is advisory. Only an authorized human may make the authoritative PASS/FAIL decision and release a pack.","warn")
doc.add_page_break()

heading(doc,"1. Purpose and release decision",1)
add_p(doc,"This guide validates SteriCord as one connected safety workflow: authenticated personnel register equipment, create and process sterilization cycles, register packs, capture evidence, confirm OCR/AI assistance, perform human verification, release eligible packs, verify them at point of use, and manage any resulting incident or quarantine with a complete audit trail and compliant reports.")
heading(doc,"1.1 Exit criteria",2)
for x in ["All P0 tests pass; no unresolved defect can permit an ineligible pack to be released or used.","All P1 tests pass, or an approved risk acceptance names an owner, mitigation, and due date.","Tenant isolation, role restrictions, append-only audit behavior, and private evidence access pass independently of UI behavior.","Every state transition produces the expected database state and audit event; OT FAIL is atomic across pack, incident, quarantine, notification, verification, and audit records.","OCR and Gemini are tested both with real credentials and through their intentional NOT_CONFIGURED manual fallback, when the pilot scope includes those services.","Mobile camera, QR, permissions, encryption, interruption, retry, and idempotency are validated on representative physical devices before clinical use.","Backup restore is proven into a clean target, and generated PDFs are readable, tenant-scoped, and audited."]:
    bullet(doc,x)
callout(doc,"Stop-ship","Any cross-tenant disclosure, unauthorized status transition, false automatic release, lost FAIL alert, duplicate offline verification, mutable audit history, or unrecoverable backup is a release blocker.","risk")

heading(doc,"2. System under test",1)
table(doc,["Layer","Coverage in this guide","Primary evidence"],[
    ("Web console","Dashboard, equipment, cycles, packs, evidence, point of use, incidents, quarantine, audit, reports, users, MFA settings","Screenshots, console/network log"),
    ("Mobile","CSSD and OT flows, camera, QR, notifications, offline queue and encrypted storage","Device video, logs, queue/database inspection"),
    ("Backend","Postgres/RLS, workflow RPCs, storage, Realtime, idempotency and retention","SQL/API results, database snapshots"),
    ("AI services","Document AI OCR and Gemini indicator classification with human confirmation","Edge logs, stored extraction/analysis rows"),
    ("Reporting","FastAPI/WeasyPrint equipment, cycle, incident and evidence PDFs","PDF, response headers, audit row"),
    ("Operations","Security checks, rate limiting, backup/restore, monitoring readiness","Command output, restored counts, alerts")],[1500,4860,3000])
heading(doc,"2.1 Authoritative state models",2)
table(doc,["Entity","Normal progression","Exception progression"],[
    ("Cycle","CREATED → RUNNING → COMPLETED → PENDING_VERIFICATION → PASSED","PENDING_VERIFICATION → FAILED; UNDETERMINED remains pending"),
    ("Pack","PACKED → STERILIZING → PENDING_VERIFICATION → RELEASED → USED","→ QUARANTINED / RECALLED / EXPIRED"),
    ("Incident","DETECTED/ALERTED → ACKNOWLEDGED → investigation states → RESOLVED/CLOSED","Terminal RESOLVED/CLOSED cannot advance"),
    ("Quarantine","ACTIVE → RELEASED","Only admin/supervisor may release with reason")],[1600,4200,3560])
callout(doc,"Implementation note","Direct client updates must not drive cycle or pack status. Workflow RPCs enforce role, tenant, current-state rules, downstream effects, and audit logging.","blue")

heading(doc,"3. Roles and test identities",1)
table(doc,["Account","Role","Core validation"],[
    ("admin@hospital-a.test","HOSPITAL_ADMIN","Users, equipment, verification, incidents, quarantine, reports"),
    ("supervisor@hospital-a.test","CSSD_SUPERVISOR","Full CSSD operate/verify flow; incident and quarantine handling"),
    ("technician@hospital-a.test","CSSD_TECHNICIAN","Create/operate cycles, packs and evidence; cannot verify cycle or release quarantine"),
    ("nurse@hospital-a.test","OT_NURSE","Point-of-use eligibility and PASS/FAIL"),
    ("admin@hospital-b.test","HOSPITAL_ADMIN","Negative tenant-isolation control")],[2700,2200,4460])
add_p(doc,"Seeded local password: SteriCord!2026. Treat it as test-only and never reuse it in a pilot environment.")
heading(doc,"3.1 Permission probes",2)
for x in ["Repeat sensitive actions through the UI, direct REST request, and RPC invocation; hiding a button alone is not authorization.","Validate that revoked and inactive roles lose access without historical grants being deleted.","Confirm anonymous and expired sessions cannot query tenant data, generate signed evidence URLs, invoke workflow functions, or generate reports.","Verify Hospital A identifiers cannot be used by Hospital B in list, detail, RPC, storage, reporting, Realtime, or guessed-URL paths."]:
    bullet(doc,x)

heading(doc,"4. Environment and data preparation",1)
for x in ["Node.js 20+, Supabase CLI, Docker Desktop; Flutter SDK plus Android/iOS tooling for mobile validation.","Start Supabase and record API URL, anon key, database port, Studio URL, migration level, Git commit, browser, device model, OS, and app build.","Copy required environment examples; never paste secrets into screenshots or the test record.","Start the web app. Start the reporting container with its service environment. Configure real Document AI/Gemini credentials only in a controlled test project.","Reset and seed the database for a repeatable run. Optionally load seed_demo_data.sql for coverage of CREATED, RUNNING, pending, passed, failed, used, and quarantined examples.","Prepare sharp, blurry, dark, glare-heavy, receipt, Bowie-Dick, chemical, biological tray, OT indicator and washer-test images. Use synthetic/non-patient identifiers.","Prepare two enrolled MFA users, two tenants, one ACTIVE and one locked/maintenance equipment record, expired and unreleased packs, and a stable client operation UUID for retry tests."]:
    bullet(doc,x)
heading(doc,"4.1 Execution record",2)
table(doc,["Field","Record before execution"],[
    ("Build / commit",""),("Environment / region",""),("Tester / date / timezone",""),("Web browser",""),("Mobile devices / OS",""),("Supabase migration head",""),("OCR / Gemini mode","Real credentials / NOT_CONFIGURED"),("Reporting image tag",""),("Evidence folder / defect tracker","" )],[2600,6760])

heading(doc,"5. Test execution method",1)
for x in ["Use a clean test identity and note the initial entity state.","Perform only the documented action; capture timestamp and visible result.","Verify UI/API response, database row(s), storage object, notification delivery, and audit event as applicable.","Refresh or sign in from a second client to prove persistence and Realtime behavior.","For negative tests, prove both that the action failed and that no partial downstream records were created.","Record Pass, Fail, Blocked, or Not Applicable. A Blocked P0 is not a pass."]:
    num(doc,x)
table(doc,["Priority","Meaning","Required disposition"],[
    ("P0","Patient safety, tenant isolation, authority, audit or recovery","Must pass before release"),("P1","Core operational/compliance capability","Pass or formal risk acceptance"),("P2","Usability, compatibility or non-critical edge","Schedule with owner")],[1000,4500,3860])

heading(doc,"6. Full-flow test cases",1)
tc(doc,"FF-01","Authentication, protected routes and session lifecycle","All roles","Backend and web are running; seeded users exist.",[
    "Open a protected URL while signed out; confirm redirect to Login.","Sign in with invalid credentials, then valid credentials.","Refresh, open a second protected route, sign out, and attempt browser Back.","Expire or revoke the session and retry a data request."],"Invalid login is generic; valid login loads the correct tenant; refresh preserves the session; logout/expiry removes protected data and blocks calls.","Screenshots and auth/network responses.")
tc(doc,"FF-02","MFA enrollment and AAL2 challenge","Admin / supervisor","A test authenticator is available and the user has no enrolled factor.",[
    "Enable two-factor authentication in Settings, scan the QR, and enter an invalid then valid TOTP.","Sign out and sign in again; complete the challenge with invalid, expired, then valid codes.","Attempt to bypass the challenge by directly opening /dashboard."],"Factor is verified only by a valid current code; next login requires AAL2; protected content is unavailable until challenge succeeds.","Enrollment state, challenge screens and auth assurance level.")
tc(doc,"FF-03","Tenant isolation across every surface","Hospital A admin and Hospital B admin","Known IDs and object paths exist in both tenants.",[
    "List and open equipment, cycles, packs, incidents, audit rows and reports as Hospital A.","Sign in as Hospital B and request Hospital A IDs through UI URL, REST, RPC, report endpoint and signed-storage URL.","Subscribe to notifications and cause a Hospital A OT failure."],"Hospital B sees no Hospital A row, object, report, alert or metadata; requests return not found/unauthorized without confirming sensitive existence.","HTTP/SQL output and Realtime capture.")
tc(doc,"FF-04","Equipment registration and operating-status control","Admin / supervisor / biomedical engineer","Hospital A user has an authorized role.",[
    "Create equipment with valid unique code and required metadata.","Attempt duplicate code and invalid/missing fields.","Change ACTIVE → MAINTENANCE → LOCKED_OUT and inspect status history and audit.","Attempt change as technician or nurse, including a direct request."],"Valid equipment persists; duplicates/invalid data fail; every status change creates history/audit; unauthorized writes fail without mutation.","Equipment row, status-event rows and audit events.")
tc(doc,"FF-05","CSSD happy path from cycle creation to pack release","Supervisor","ACTIVE equipment and instrument set exist.",[
    "Create a cycle and register two uniquely coded packs.","Start the cycle; confirm cycle RUNNING and packs STERILIZING.","Complete with temperature, pressure and hold time; upload accepted evidence.","Submit for verification; record human PASS; release each pack."],"States progress only in order; timestamps are populated once; evidence is private; PASS changes cycle to PASSED; only its pending packs can become RELEASED; audit records identify actor/action/entity/time.","Before/after UI, database rows, storage metadata and audit chain.")
tc(doc,"FF-06","CSSD state-machine and authorization negatives","Technician / supervisor","Cycles and packs exist in representative states.",[
    "Try to complete CREATED, start RUNNING, submit RUNNING, verify COMPLETED, or release a pack before cycle PASS.","Try to register a pack after the cycle is past RUNNING.","As technician, attempt cycle verification and quarantine release via UI and RPC.","Attempt direct status UPDATE through REST."],"Every invalid transition/role/direct update is rejected and produces no partial verification, release, or status change.","Error responses plus row-count/state comparison.")
tc(doc,"FF-07","Cycle FAIL and UNDETERMINED routing","Supervisor","Two PENDING_VERIFICATION cycles with packs exist.",[
    "Verify the first cycle FAIL with notes.","Verify the second UNDETERMINED.","Inspect all related packs, verifications and audit records.","Attempt release from each outcome."],"FAIL sets cycle FAILED and pending packs QUARANTINED; UNDETERMINED records review but keeps cycle pending; neither path permits release.","Cycle/pack snapshots and verification/audit rows.")
tc(doc,"FF-08","Evidence quality gate and private storage","Technician / nurse","Prepared sharp, blurry, dark and glare-heavy images exist.",[
    "For each relevant evidence type, upload a sharp image and record quality scores.","Attempt blurry, dark and glare-heavy images; confirm specific retake reasons.","Verify rejected images create neither storage objects nor evidence rows.","Attempt cross-tenant/object-path access and allow a signed URL to expire."],"Accepted images persist with type-specific metrics and seven-year retention date; rejected images never upload; private objects require authorized short-lived URLs.","Image set, UI messages, storage listing and evidence rows.")
tc(doc,"FF-09","OCR extraction with human confirmation","Technician / supervisor","Accepted AUTOCLAVE_RECEIPT evidence exists.",[
    "Run OCR with real credentials and inspect extracted fields/confidence.","Correct at least one value and confirm the human fields.","Repeat with service deliberately unconfigured and enter values manually.","Attempt confirmation from another tenant or unauthorized role."],"Real OCR pre-fills but never becomes authoritative until confirmation; edits are preserved separately; NOT_CONFIGURED gives a clear manual path; unauthorized confirmation fails.","Edge logs and ocr_extractions row.","P1")
tc(doc,"FF-10","Gemini advisory analysis and confidence routing","Supervisor / nurse","Accepted supported indicator images exist.",[
    "Analyze representative PASS, FAIL and ambiguous images.","Confirm HIGH ≥ 0.85, MEDIUM ≥ 0.60, and LOW < 0.60 bands.","Verify LOW is stored/presented as UNDETERMINED, regardless of raw model wording.","Choose a human result that differs from the suggestion and submit."],"AI remains visibly advisory, never auto-submits or releases; confidence is bounded; low confidence requires manual review; verification stores both AI suggestion/confidence and independent human result.","AI rows, UI banner and verification row.","P1")
tc(doc,"FF-11","OT point-of-use PASS","OT nurse","A non-expired RELEASED pack exists.",[
    "Scan the QR or enter the pack code.","Confirm eligibility and pack identity before recording PASS.","Record PASS with synthetic encounter references and notes.","Attempt to scan the same pack again."],"One point-of-use record and verification are created; pack becomes USED; audit event is present; subsequent use is ineligible and creates no duplicate record.","Scan capture and linked database rows.")
tc(doc,"FF-12","OT point-of-use eligibility rejection","OT nurse","Prepare PACKED, QUARANTINED, USED, EXPIRED and unknown pack codes.",[
    "Scan each ineligible/unknown code.","Try to bypass the displayed rejection by calling the verification RPC.","Confirm no incident is created merely by an eligibility rejection."],"Each is blocked with a useful non-sensitive reason; server independently enforces eligibility; no verification/status/incident side effects occur.","UI/API messages and zero-delta row counts.")
tc(doc,"FF-13","OT FAIL atomic safety cascade","OT nurse; observing supervisor","A RELEASED eligible pack exists; supervisor has notifications open.",[
    "Record a FAIL with notes.","Immediately refresh pack, incident, quarantine, notifications, verification and audit views.","Confirm the supervisor receives a CRITICAL Realtime notification and marks it read.","Simulate one invalid downstream condition and prove the transaction does not partially commit."],"In one transaction the pack becomes QUARANTINED; point-of-use verification, ALERTED CRITICAL incident, ACTIVE quarantine, notifications and audit events exist. Failure rolls back the entire cascade.","Correlated IDs/timestamps across all records and Realtime event.")
tc(doc,"FF-14","Incident acknowledgement, progression and closure","Admin / supervisor","An ALERTED incident from FF-13 exists.",[
    "Acknowledge the incident.","Advance through an appropriate investigation/replacement state.","Resolve with mandatory notes, then test close behavior.","Attempt regression, advance from a terminal status, and actions as nurse."],"Only authorized roles advance; history is coherent; terminal records cannot advance; resolution metadata is retained and audited; unauthorized attempts fail.","Incident timeline and audit events.")
tc(doc,"FF-15","Quarantine release","Admin / supervisor","ACTIVE quarantine and related QUARANTINED pack exist.",[
    "Attempt release without a reason and as an unauthorized role.","Release with a clinically approved reason.","Inspect quarantine and linked pack state; attempt a second release.","Confirm incident resolution and quarantine release are distinct actions."],"Only authorized, reasoned release succeeds once; quarantine becomes RELEASED; pack is not silently returned to service unless explicitly defined by the workflow; all activity is audited.","Quarantine/pack rows and audit events.")
tc(doc,"FF-16","Offline OT queue, encryption and idempotent retry","OT nurse on mobile","Physical device, released pack, network controls and fixed client_operation_id.",[
    "Go offline, record point-of-use PASS/FAIL, and inspect that queued payload is encrypted at rest.","Restart the app and confirm the queued operation remains recoverable.","Reconnect, interrupt during sync, then retry the same operation ID multiple times.","Confirm a failed item retries at a bounded rate and does not block later eligible items."],"Queue states progress safely; interrupted SYNCING recovers; server returns the original result for the same operation ID; exactly one point-of-use record and one safety cascade exist.","Device video/logs, encrypted database inspection and server row counts.")
tc(doc,"FF-17","Realtime notification isolation and read state","Admin / supervisor","Two tenants and two logged-in recipients exist.",[
    "Trigger a Hospital A OT FAIL.","Observe authorized Hospital A recipients and Hospital B client.","Mark one notification read and reload.","Reconnect after a brief network loss."],"Only intended active roles in Hospital A receive the CRITICAL alert; read state persists per notification; reconnect does not leak or multiply notifications.","Realtime logs and notification rows.","P1")
tc(doc,"FF-18","Audit chain integrity and immutability","Admin / auditor","Multiple critical events exist.",[
    "Filter the audit trail and verify actor, event, entity, before/after state and timestamp.","Run Verify Chain Integrity.","Attempt client INSERT, UPDATE and DELETE on audit_events.","In an isolated disposable environment, tamper as database owner and rerun integrity verification."],"Normal clients cannot alter history; intact chain passes; controlled tamper identifies the affected row/chain failure; tenant scope is enforced.","Audit UI, RPC result and database error/output.")
tc(doc,"FF-19","Compliance report generation","Authorized role","Reporting container runs and equipment/cycle/incident/evidence IDs exist.",[
    "Generate each of the four PDF types.","Validate title, identifiers, state, timestamps, evidence/verification linkage and readable pagination.","Attempt wrong-tenant, unknown and unauthorized IDs.","Generate repeatedly beyond 30 requests/minute for one user, then test another user."],"Authorized PDFs contain only permitted data and each export creates REPORT_EXPORTED audit evidence; invalid access fails; per-user rate limit returns 429 without affecting other users.","PDFs, response codes, service logs and audit rows.","P1")
tc(doc,"FF-20","User and role lifecycle","Hospital admin","A new test user and multiple roles are available.",[
    "Create/activate a profile and grant an allowed non-system role.","Verify permissions, grant a second role, revoke the first, and deactivate the user.","Attempt duplicate active grant and SYSTEM_ADMIN assignment.","Confirm historical grant/revocation fields remain available."],"Effective access follows active grants; prohibited/duplicate grants fail; deactivation blocks access; history is preserved rather than deleted.","User-role rows and authorization probes.","P1")
tc(doc,"FF-21","Backup, restore and continuity","Operations","A populated environment and isolated restore target exist.",[
    "Record source row counts and representative hashes/IDs.","Run the backup script and verify a non-empty protected artifact.","Restore into the clean target using the restore script.","Run key queries, audit-chain verification and one read-only application smoke test."],"Restore completes without overwriting the source; schema/data counts and sampled records match; audit chain remains intact; the restored app can read expected tenant data.","Backup metadata, command logs, comparisons and smoke-test result.")
tc(doc,"FF-22","Mobile device and camera acceptance","Technician / OT nurse","Representative Android and iOS devices are available.",[
    "Install cleanly; test denied/allowed camera permission, rotation, background/foreground and low-storage behavior.","Scan printed and screen QR codes at varied distance, angle and OT lighting.","Capture sharp/blurry/dark/glare images and compare the quality gate across devices.","Exercise login, MFA, cycle, evidence, point-of-use, notifications and offline flows."],"No crash or stuck camera; permission recovery is clear; QR and capture are reliable enough for the pilot environment; quality thresholds do not show unacceptable device-specific false results.","Device matrix, videos, timings and defect IDs.")

heading(doc,"7. Cross-cutting non-functional checks",1)
table(doc,["Area","Minimum checks","Release signal"],[
    ("Accessibility","Keyboard path, visible focus, labels, error association, contrast, zoom/text scaling, screen reader basics","Critical workflows usable without pointer and errors announced"),
    ("Performance","Dashboard/list/detail latency, upload/analysis time, report latency, Realtime delivery, sync backlog","Pilot targets agreed and met at representative concurrency"),
    ("Reliability","Refresh during mutation, duplicate clicks, timeout after commit, reconnect, service restart","No duplicate/partial safety record; recovery is clear"),
    ("Privacy","No patient data in URLs/logs/errors; signed URLs expire; screenshots use synthetic data","No unauthorized persistence or disclosure"),
    ("Compatibility","Supported browsers; representative iOS/Android; camera and file formats","Documented support matrix passes"),
    ("Observability","Structured logs, correlation IDs, health checks, alert routing, clock sync","On-call can detect and trace a failed critical flow")],[1500,4860,3000])

heading(doc,"8. Data integrity reconciliation",1)
for x in ["Each cycle verification has one authoritative human result and correct verifier/timestamp.","Released packs belong only to PASSED cycles; USED packs have a successful point-of-use record.","Every OT FAIL has a linked FAIL verification, incident, ACTIVE/managed quarantine, CRITICAL notification fan-out, and required audit events.","Every accepted evidence row maps to one private object; rejected captures map to neither.","Every OCR confirmation preserves machine extraction and separate human-confirmed fields.","Every AI-assisted verification preserves suggestion/confidence and independent human result; LOW confidence never remains authoritative PASS/FAIL.","Every exported report has an audit event; retention_expires_at is approximately seven years for evidence and audit rows.","No duplicate client_operation_id exists per organization and retry returns the original result_entity_id."]:
    bullet(doc,x)

heading(doc,"9. Defect triage and evidence package",1)
table(doc,["Severity","Examples","Immediate action"],[
    ("S0 / Critical","Cross-tenant disclosure; ineligible release/use; missing OT FAIL cascade; audit mutation; unrecoverable backup","Stop testing/release, preserve evidence, notify safety/security owner"),
    ("S1 / High","Authorized core flow unavailable; incorrect status; missing report/audit/notification","Block release until fixed or formally accepted"),
    ("S2 / Medium","Recoverable workflow friction, compatibility or misleading message","Fix before pilot where operationally material"),
    ("S3 / Low","Cosmetic or minor copy issue","Backlog with evidence")],[1500,5100,2760])
for x in ["Test ID, build/commit, environment, exact role and tenant.","Pre-state IDs/statuses and sanitized test data.","Timestamp with timezone, steps, expected and actual result.","Screenshots/video, request/response, device/browser logs and relevant SQL output.","Downstream row counts and correlation IDs proving whether the transaction partially committed.","Safety/privacy impact, reproducibility, workaround, owner and retest result."]:
    bullet(doc,x)

heading(doc,"10. Pilot-specific validation gaps",1)
callout(doc,"Current baseline","The repository reports web/server phases as locally verified, but real Google OCR/Gemini credentials, Flutter compilation, physical device/camera/QR behavior, production monitoring, scheduled off-host backups, and pilot-scale load remain environment-dependent validations.","warn")
for x in ["Run flutter pub get, flutter analyze, unit/widget tests and release builds before device testing.","Test killed-app recovery from SYNCING and duplicate retry on real devices; server idempotency alone does not prove client correctness.","Confirm CSSD Wi-Fi assumptions. Offline coverage currently targets OT point of use, not cycle operations.","Tune image-quality thresholds using the actual hospital devices, lighting and evidence media; document false-accept/false-reject results.","Wire logs to an approved monitoring/alerting platform and conduct an alert drill.","Schedule backups to separate storage and perform a timed recovery rehearsal.","Obtain compliance sign-off that retention is recorded but automated archival/deletion enforcement is not yet implemented."]:
    bullet(doc,x)

heading(doc,"Appendix A — Test run summary",1)
table(doc,["ID","Result","Defect / evidence reference","Tester / date"],[(f"FF-{i:02d}","□ Pass  □ Fail  □ Blocked  □ N/A","","") for i in range(1,23)],[900,2300,4200,1960])

heading(doc,"Appendix B — Release sign-off",1)
table(doc,["Gate","Decision / evidence"],[
    ("P0 complete",""),("P1 disposition complete",""),("Security and tenant isolation",""),("Clinical workflow SME approval",""),("Mobile/device acceptance",""),("Backup/restore rehearsal",""),("Monitoring and support readiness",""),("Final release decision","GO / NO-GO" )],[3500,5860])
add_p(doc,"Approvals")
table(doc,["Role","Name","Signature / date"],[("QA lead","",""),("CSSD clinical lead","",""),("OT clinical lead","",""),("Security/privacy","",""),("Product owner","","")],[2600,3000,3760])

# Keep tables from splitting rows and add repeated headers.
for t in doc.tables:
    if t.rows:
        trPr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement("w:tblHeader"); rep.set(qn("w:val"),"true"); trPr.append(rep)
    for row in t.rows:
        pr=row._tr.get_or_add_trPr(); cant=OxmlElement("w:cantSplit"); pr.append(cant)

OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)
